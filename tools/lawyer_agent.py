#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
AI-агент адвоката (бот «Судові засідання та практика», кнопка «🤖 Завдання агенту»).

Адвокат надсилає документи (Word DOC/DOCX, PDF, фото; один чи кілька) і ставить
завдання своїми словами або кнопкою-шаблоном: аналіз, відзив, заперечення,
апеляційна скарга, клопотання тощо. Агент сам визначає форму результату:
  • документ  — DOCX (оформлення за ДСТУ 4163-2020) + перелік «що перевірити»;
  • відповідь — розбір/висновок текстом у чат.

Завдання воркфлоу отримує лише як id (AGENT_JOB); текст і file_id забирає
з Worker (/agent_job, одноразово). Репозиторій публічний, логи відкриті — тому
сюди НЕ друкуємо ні тексту справи, ні ПІБ, ні назв файлів: лише лічильники.

Env: AGENT_JOB, BOT_WORKER_URL, BOT_WORKER_SECRET, TELEGRAM_BOT_TOKEN,
     TELEGRAM_REVIEW_CHAT, ANTHROPIC_API_KEY, AGENT_MODEL (необов'язково).
"""
import base64
import io
import json
import os
import re
import subprocess
import sys
import time
import urllib.parse
import urllib.request
import uuid

TOKEN = os.environ.get("TELEGRAM_BOT_TOKEN", "").strip()
CHAT = os.environ.get("TELEGRAM_REVIEW_CHAT", "").strip()
WORKER_URL = os.environ.get("BOT_WORKER_URL", "").strip().rstrip("/")
WORKER_SECRET = os.environ.get("BOT_WORKER_SECRET", "").strip()
MODEL = os.environ.get("AGENT_MODEL", "claude-opus-5-5").strip()
ADVOCATE = "Осадько Олександр Олексійович"

SEP = "===ПЕРЕВІРИТИ==="
MODE_DOC = "РЕЖИМ: ДОКУМЕНТ"
MODE_ANS = "РЕЖИМ: ВІДПОВІДЬ"

SYSTEM = f"""Ти — досвідчений помічник українського адвоката ({ADVOCATE}). Адвокат
надсилає документи справи (рішення, ухвали, позови, скарги, протоколи, договори тощо)
і ставить завдання. Виконай його українською мовою за чинним законодавством України
(ЦПК, ГПК, КАС, КПК, КУпАП, ЦК, СК, КЗпП тощо), спираючись на зміст наданих документів.

Спершу визнач форму результату і ПЕРШИМ рядком відповіді напиши рівно одне:
  {MODE_DOC}   — якщо треба скласти процесуальний чи інший документ (відзив,
                 заперечення, позов, апеляційна/касаційна скарга, клопотання, заява,
                 пояснення, лист, претензія тощо);
  {MODE_ANS}  — якщо треба проаналізувати, відповісти на питання, оцінити перспективи.

Загальні правила:
- Посилайся на конкретні статті (частини, пункти) законів. Позиції Верховного Суду
  наводь лише ті, у яких упевнений (номер справи і дата); якщо не впевнений — не
  вигадуй, а напиши «усталена практика ВС» і внеси це в перелік перевірки.
- Факти бери з документів; невідомі дані (ПІБ, адреси, РНОКПП, номери, дати, суми)
  не вигадуй — став місце для заповнення у квадратних дужках: [ПІБ позивача].
- Представник — адвокат {ADVOCATE}; реквізити ордера/договору — у дужках.
- Враховуй строки: для документа-відповіді зазнач строк його подання, для рішень —
  строк і порядок оскарження (з якого моменту рахується).

Якщо {MODE_DOC}: далі — ЛИШЕ текст документа простою розміткою, без markdown-зірочок:
  >> рядок         — шапка праворуч (суд, сторони, адреси, номер справи), по рядку
  # НАЗВА          — назва документа (по центру, великими)
  ## Підзаголовок  — назва розділу (по центру)
  - пункт          — пункт переліку («ПРОШУ:», «Додатки:»)
  звичайний рядок  — абзац (по ширині, з відступом)
  _[дата] | О. О. Осадько_ — останній рядок: дата ліворуч, підпис праворуч
Документ — повний, аргументований, з мотивувальною частиною, що спростовує доводи
іншої сторони по пунктах. Після документа окремим рядком напиши {SEP}, а далі —
3–8 пунктів (кожен з «- »): що перевірити чи дописати перед поданням (строки,
судовий збір, докази, сумнівні посилання).

Якщо {MODE_ANS}: стислий практичний розбір без вступів. Для аналізу рішення чи
позову — за схемою: <b>📄 Що це</b>, <b>⚖️ Суть</b>, <b>🧾 Що вирішено / чого просять</b>,
<b>⏰ Строки</b> (норма, з якого моменту, кінцева дата, якщо видно дату),
<b>🎯 Слабкі місця і аргументи</b>, <b>✅ Що зробити</b>. Якщо поставлено питання —
спершу коротка відповідь на нього. Формат: Telegram HTML — лише теги <b> та <i>,
без markdown і таблиць."""


def log(msg):
    print(msg, flush=True)


# ─────────────────────────── Worker і Telegram ───────────────────────────
def fetch_job(job_id):
    req = urllib.request.Request(
        f"{WORKER_URL}/agent_job?id={urllib.parse.quote(job_id)}",
        headers={"X-Auth-Token": WORKER_SECRET, "User-Agent": "osadko-agent"})
    with urllib.request.urlopen(req, timeout=30) as r:
        return json.loads(r.read().decode("utf-8"))


def tg_api(method, data=None, files=None):
    url = f"https://api.telegram.org/bot{TOKEN}/{method}"
    if files:
        boundary = uuid.uuid4().hex
        body = io.BytesIO()
        for k, v in (data or {}).items():
            body.write(f"--{boundary}\r\nContent-Disposition: form-data; name=\"{k}\"\r\n\r\n"
                       f"{v}\r\n".encode("utf-8"))
        for k, (fname, content, ctype) in files.items():
            body.write(f"--{boundary}\r\nContent-Disposition: form-data; name=\"{k}\"; "
                       f"filename=\"{fname}\"\r\nContent-Type: {ctype}\r\n\r\n".encode("utf-8"))
            body.write(content)
            body.write(b"\r\n")
        body.write(f"--{boundary}--\r\n".encode())
        req = urllib.request.Request(url, data=body.getvalue(), headers={
            "Content-Type": f"multipart/form-data; boundary={boundary}"})
    else:
        req = urllib.request.Request(url, data=json.dumps(data or {}).encode("utf-8"),
                                     headers={"Content-Type": "application/json"})
    with urllib.request.urlopen(req, timeout=60) as r:
        return json.loads(r.read().decode("utf-8"))


def tg_text(text):
    """Надсилає HTML-текст частинами до ~3900 символів (ліміт Telegram 4096)."""
    parts, cur = [], ""
    for para in text.split("\n"):
        if len(cur) + len(para) + 1 > 3900 and cur:
            parts.append(cur)
            cur = ""
        cur += para + "\n"
    if cur.strip():
        parts.append(cur)
    for p in parts:
        try:
            r = tg_api("sendMessage", {"chat_id": CHAT, "text": p, "parse_mode": "HTML",
                                       "disable_web_page_preview": True})
            if not r.get("ok"):
                raise RuntimeError("html")
        except Exception:  # зламаний HTML від моделі — надсилаємо як звичайний текст
            tg_api("sendMessage", {"chat_id": CHAT, "text": re.sub(r"<[^>]+>", "", p)})
        time.sleep(0.4)


def tg_download(file_id):
    info = tg_api("getFile", {"file_id": file_id})
    path = info["result"]["file_path"]
    with urllib.request.urlopen(f"https://api.telegram.org/file/bot{TOKEN}/{path}",
                                timeout=60) as r:
        return r.read(), path


# ───────────────────────────── вміст для Claude ─────────────────────────────
def doc_to_text(raw):
    """Старий Word (.doc) → текст через antiword (ставимо за потреби)."""
    import shutil
    import tempfile
    if not shutil.which("antiword"):
        subprocess.run(["sudo", "apt-get", "install", "-y", "-qq", "antiword"],
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, check=False)
    with tempfile.NamedTemporaryFile(suffix=".doc") as f:
        f.write(raw)
        f.flush()
        try:
            r = subprocess.run(["antiword", "-m", "UTF-8.txt", "-w", "0", f.name],
                               capture_output=True, timeout=60)
        except (OSError, subprocess.TimeoutExpired):
            raise RuntimeError("не вдалося прочитати .doc — збережіть як DOCX або PDF")
    return r.stdout.decode("utf-8", "replace")


def file_blocks(file, n):
    """Файл із Telegram → блоки повідомлення для Claude (PDF / фото / Word / текст)."""
    raw, path = tg_download(file["file_id"])
    name = (file.get("name") or path).lower()
    mime = (file.get("mime") or "").lower()
    ext = name.rsplit(".", 1)[-1] if "." in name else mime
    log(f"Файл {n}: {len(raw) // 1024} КБ, тип {ext}")
    label = {"type": "text", "text": f"Документ {n} ({file.get('name') or 'файл'}):"}
    if name.endswith(".pdf") or mime == "application/pdf":
        return [label, {"type": "document", "source": {
            "type": "base64", "media_type": "application/pdf",
            "data": base64.b64encode(raw).decode()}}]
    if mime.startswith("image/") or re.search(r"\.(jpe?g|png|webp|gif)$", name):
        mt = mime if mime in ("image/jpeg", "image/png", "image/webp", "image/gif") else "image/jpeg"
        return [label, {"type": "image", "source": {
            "type": "base64", "media_type": mt, "data": base64.b64encode(raw).decode()}}]
    if name.endswith(".docx") or "wordprocessingml" in mime:
        import docx
        d = docx.Document(io.BytesIO(raw))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        text = "\n".join(parts)
    elif name.endswith(".doc") or mime == "application/msword":
        text = doc_to_text(raw)
    else:
        try:
            text = raw.decode("utf-8")
        except UnicodeDecodeError:
            raise RuntimeError("непідтримуваний тип файлу")
    if not text.strip():
        raise RuntimeError("не вдалося прочитати текст файлу")
    return [{"type": "text", "text": f"Документ {n} ({file.get('name') or 'файл'}), текст:\n\n"
             + text[:150000]}]


def ask_claude(system, content, max_tokens):
    import anthropic
    client = anthropic.Anthropic()
    r = client.beta.messages.create(
        model=MODEL,
        max_tokens=max_tokens,
        output_config={"effort": "high"},
        betas=["server-side-fallback-2026-07-01"],
        fallbacks="default",
        system=system,
        messages=[{"role": "user", "content": content}],
    )
    if r.stop_reason == "refusal":
        raise RuntimeError("модель відмовилась відповідати")
    log(f"Claude: вхід {r.usage.input_tokens} · вихід {r.usage.output_tokens} токенів")
    return "".join(b.text for b in r.content if getattr(b, "type", "") == "text").strip()


# ─────────────────────────── DOCX за ДСТУ 4163-2020 ───────────────────────────
def build_docx(text):
    """Проста розмітка моделі → DOCX: Times New Roman 14, поля 30/10/20/20 мм,
    абзацний відступ 1,25 см, вирівнювання по ширині, міжрядковий 1,15."""
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    d = docx.Document()
    sec = d.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21)
    sec.left_margin, sec.right_margin = Cm(3), Cm(1)
    sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)
    st = d.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(14)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    pf = st.paragraph_format
    pf.space_after = Pt(0)
    pf.space_before = Pt(0)
    pf.line_spacing = 1.15

    def para(txt, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True, bold=False,
             left=None, before=0, upper=False):
        p = d.add_paragraph()
        p.alignment = align
        if indent:
            p.paragraph_format.first_line_indent = Cm(1.25)
        if left is not None:
            p.paragraph_format.left_indent = left
        if before:
            p.paragraph_format.space_before = Pt(before)
        run = p.add_run(txt.upper() if upper else txt)
        run.bold = bold
        return p

    for raw in text.splitlines():
        line = raw.rstrip()
        s = line.strip().replace("**", "")
        if not s:
            continue
        if s.startswith(">>"):
            para(s[2:].strip(), WD_ALIGN_PARAGRAPH.LEFT, indent=False, left=Cm(8.5))
        elif s.startswith("## "):
            para(s[3:].strip(), WD_ALIGN_PARAGRAPH.CENTER, indent=False, bold=True, before=6)
        elif s.startswith("# "):
            para(s[2:].strip(), WD_ALIGN_PARAGRAPH.CENTER, indent=False, bold=True,
                 before=12, upper=True)
        elif s.startswith("- "):
            para("– " + s[2:].strip(), indent=True)
        elif s.startswith("_") and s.endswith("_") and len(s) > 2:
            # дата ліворуч, «________ підпис» праворуч (табуляція по правому полю)
            left, _, right = s[1:-1].partition("|")
            p = para(left.strip(" _"), WD_ALIGN_PARAGRAPH.LEFT, indent=False, before=18)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(17), WD_TAB_ALIGNMENT.RIGHT)
            if right.strip(" _"):
                p.add_run("\t____________ " + right.strip(" _"))
        else:
            para(s)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def doc_title(text):
    for line in text.splitlines():
        if line.strip().startswith("# "):
            return line.strip()[2:].strip()
    return "Документ"


def safe_filename(title):
    t = re.sub(r"[^\w\s-]", "", title, flags=re.UNICODE).strip()
    t = re.sub(r"\s+", "_", t)[:60] or "Документ"
    return f"Чернетка_{t}_{time.strftime('%d.%m.%Y')}.docx"


# ─────────────────────────────────── main ───────────────────────────────────
def send_document(doc, check):
    title = doc_title(doc)
    data = build_docx(doc)
    cap = f"📝 <b>{title}</b> — чернетка.\nПеревірте й доповніть перед поданням."
    tg_api("sendDocument", {"chat_id": CHAT, "caption": cap[:1000], "parse_mode": "HTML"},
           files={"document": (safe_filename(title), data,
                               "application/vnd.openxmlformats-officedocument."
                               "wordprocessingml.document")})
    items = [re.sub(r"^[-•]\s*", "", x).strip() for x in check.splitlines() if x.strip()]
    if items:
        esc = lambda x: x.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        tg_text("🔍 <b>Що перевірити:</b>\n" + "\n".join(f"• {esc(x)}" for x in items))
    log(f"Документ: {len(doc)} символів, пунктів перевірки: {len(items)}")


def run(job):
    files = job.get("files") or []
    content = []
    for i, f in enumerate(files, 1):
        content += file_blocks(f, i)
    task = (job.get("text") or "").strip() or "Проаналізуй додані документи."
    content.append({"type": "text", "text": "Завдання адвоката:\n" + task})
    out = ask_claude(SYSTEM, content, 16000)
    first, _, rest = out.partition("\n")
    if MODE_DOC in first:
        doc, _, check = rest.partition(SEP)
        send_document(doc, check)
    else:
        text = rest if MODE_ANS in first else out
        tg_text(text.strip())
        log(f"Відповідь: {len(text)} символів")


def main():
    job_id = os.environ.get("AGENT_JOB", "").strip()
    if not (job_id and WORKER_URL and WORKER_SECRET and TOKEN and CHAT):
        log("Немає AGENT_JOB або секретів воркера/Telegram.")
        sys.exit(1)
    try:
        job = fetch_job(job_id)
    except Exception as e:
        log(f"Завдання не знайдено або вже виконане: {type(e).__name__}")
        sys.exit(1)
    log(f"Завдання: {len(job.get('text') or '')} символів · файлів: {len(job.get('files') or [])}")
    try:
        run(job)
    except Exception as e:
        # У лог — лише тип помилки (повідомлення може містити фрагменти даних).
        log(f"Помилка: {type(e).__name__}")
        reason = str(e) if isinstance(e, RuntimeError) else type(e).__name__
        tg_api("sendMessage", {"chat_id": CHAT, "text":
                               f"⚠️ Не вдалося виконати завдання ({reason}). "
                               "Спробуйте ще раз або надішліть файл в іншому форматі (PDF чи DOCX)."})
        sys.exit(1)


if __name__ == "__main__":
    main()
